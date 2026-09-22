"""
Creates a sample Excel workbook users can fill with their client data.
Run directly:  python generate_sample.py
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ── palette ────────────────────────────────────────────────────────────────────
BLUE = "1E3A5F"
LIGHT_BLUE = "2E86AB"
ACCENT = "EAF4FB"
WHITE = "FFFFFF"
YELLOW = "FFF9C4"
THIN = Side(style="thin", color="BDC3C7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _hdr(ws, row: int, col: int, value: str, bg: str = BLUE, fg: str = WHITE):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name="Arial", bold=True, color=fg, size=10)
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER
    return cell


def _data(ws, row: int, col: int, value, bg: str = WHITE, bold: bool = False):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name="Arial", bold=bold, size=9)
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(vertical="center")
    cell.border = BORDER
    return cell


CLIENTS = [
    {
        "account_no": "ACC001",
        "name": "Rahul Sharma",
        "email": "rahul.sharma@example.com",
        "phone": "9876543210",
        "address": "12 MG Road, Mumbai 400001",
        "account_type": "Savings",
        "opening_balance": 50000.00,
    },
    {
        "account_no": "ACC002",
        "name": "Priya Mehta",
        "email": "priya.mehta@example.com",
        "phone": "9812345678",
        "address": "45 Park Street, Kolkata 700016",
        "account_type": "Current",
        "opening_balance": 125000.00,
    },
    {
        "account_no": "ACC003",
        "name": "Arjun Reddy",
        "email": "arjun.reddy@example.com",
        "phone": "9000123456",
        "address": "7 Jubilee Hills, Hyderabad 500033",
        "account_type": "Savings",
        "opening_balance": 80000.00,
    },
]

TRANSACTIONS = [
    # ACC001
    ("ACC001", "01-Aug-2026", "Opening Balance",        "",          "",        0, 50000.00),
    ("ACC001", "05-Aug-2026", "UPI – Amazon Pay",        "UPI/87654", 3500.00,  0, 46500.00),
    ("ACC001", "10-Aug-2026", "Salary Credit",           "NEFT/12345",0, 60000.00, 106500.00),
    ("ACC001", "14-Aug-2026", "Electricity Bill",        "NACH/0021", 2100.00,  0, 104400.00),
    ("ACC001", "18-Aug-2026", "UPI – Swiggy",            "UPI/99001", 850.00,   0, 103550.00),
    ("ACC001", "22-Aug-2026", "Rent Transfer",           "NEFT/55321",15000.00, 0,  88550.00),
    ("ACC001", "28-Aug-2026", "Interest Credit",         "",          0,  420.50,  88970.50),
    # ACC002
    ("ACC002", "01-Aug-2026", "Opening Balance",         "",          0,         0, 125000.00),
    ("ACC002", "03-Aug-2026", "Client Payment Received", "NEFT/67890",0, 45000.00, 170000.00),
    ("ACC002", "07-Aug-2026", "Supplier Payment",        "RTGS/34567",30000.00, 0, 140000.00),
    ("ACC002", "12-Aug-2026", "GST Payment",             "TDS/00123", 12500.00, 0, 127500.00),
    ("ACC002", "15-Aug-2026", "Client Payment Received", "NEFT/78901",0, 20000.00, 147500.00),
    ("ACC002", "25-Aug-2026", "Office Rent",             "NACH/4411", 22000.00, 0, 125500.00),
    # ACC003
    ("ACC003", "01-Aug-2026", "Opening Balance",         "",          0,         0,  80000.00),
    ("ACC003", "04-Aug-2026", "SIP Debit",               "NACH/SIP01",5000.00,  0,  75000.00),
    ("ACC003", "08-Aug-2026", "Freelance Income",        "IMPS/33210",0, 25000.00, 100000.00),
    ("ACC003", "11-Aug-2026", "Car EMI",                 "NACH/CAR12",8500.00,  0,  91500.00),
    ("ACC003", "18-Aug-2026", "Grocery – DMart",         "UPI/56321", 4200.00,  0,  87300.00),
    ("ACC003", "27-Aug-2026", "Dividend Credit",         "DIV/LTI22", 0,  1200.00, 88500.00),
]


def make_sample_excel(out_path: Path) -> Path:
    out_path = Path(out_path)
    wb = openpyxl.Workbook()

    # ── Clients sheet ──────────────────────────────────────────────────────────
    ws_c = wb.active
    ws_c.title = "Clients"
    ws_c.sheet_view.showGridLines = False
    ws_c.row_dimensions[1].height = 30

    c_headers = ["Account No", "Client Name", "Email", "Phone", "Address", "Account Type", "Opening Balance"]
    c_widths   = [14, 22, 28, 14, 36, 14, 16]
    for col, (h, w) in enumerate(zip(c_headers, c_widths), 1):
        _hdr(ws_c, 1, col, h)
        ws_c.column_dimensions[chr(64 + col)].width = w

    for r, client in enumerate(CLIENTS, 2):
        bg = ACCENT if r % 2 == 0 else WHITE
        _data(ws_c, r, 1, client["account_no"],    bg, bold=True)
        _data(ws_c, r, 2, client["name"],          bg)
        _data(ws_c, r, 3, client["email"],         bg)
        _data(ws_c, r, 4, client["phone"],         bg)
        _data(ws_c, r, 5, client["address"],       bg)
        _data(ws_c, r, 6, client["account_type"],  bg)
        cell = _data(ws_c, r, 7, client["opening_balance"], bg)
        cell.number_format = '#,##0.00'

    ws_c.freeze_panes = "A2"

    # ── Transactions sheet ────────────────────────────────────────────────────
    ws_t = wb.create_sheet("Transactions")
    ws_t.sheet_view.showGridLines = False
    ws_t.row_dimensions[1].height = 30

    t_headers = ["Account No", "Date", "Description", "Reference", "Debit", "Credit", "Balance"]
    t_widths   = [14, 14, 32, 18, 14, 14, 16]
    for col, (h, w) in enumerate(zip(t_headers, t_widths), 1):
        _hdr(ws_t, 1, col, h)
        ws_t.column_dimensions[chr(64 + col)].width = w

    for r, txn in enumerate(TRANSACTIONS, 2):
        bg = ACCENT if r % 2 == 0 else WHITE
        acct, date, desc, ref, deb, cred, bal = txn
        _data(ws_t, r, 1, acct, bg, bold=True)
        _data(ws_t, r, 2, date, bg)
        _data(ws_t, r, 3, desc, bg)
        _data(ws_t, r, 4, ref,  bg)
        for col_idx, val in [(5, deb), (6, cred), (7, bal)]:
            cell = _data(ws_t, r, col_idx, val if val else None, bg)
            cell.number_format = '#,##0.00'
            cell.alignment = Alignment(horizontal="right", vertical="center")

    ws_t.freeze_panes = "A2"

    # ── Instructions sheet ────────────────────────────────────────────────────
    ws_i = wb.create_sheet("Instructions")
    ws_i.sheet_view.showGridLines = False
    ws_i.column_dimensions["A"].width = 5
    ws_i.column_dimensions["B"].width = 25
    ws_i.column_dimensions["C"].width = 55

    rows = [
        ("FINTECH SAAS — DATA ENTRY GUIDE", "", BLUE),
        ("", "", WHITE),
        ("CLIENTS SHEET", "", LIGHT_BLUE),
        ("Column", "What to enter", BLUE),
        ("Account No", "Unique account number (e.g. ACC001)", WHITE),
        ("Client Name", "Full name of the client", ACCENT),
        ("Email", "Client's email address for statement delivery", WHITE),
        ("Phone", "10-digit mobile number", ACCENT),
        ("Address", "Full mailing address", WHITE),
        ("Account Type", "Savings / Current / Loan / Demat etc.", ACCENT),
        ("Opening Balance", "Balance at start of statement period", WHITE),
        ("", "", WHITE),
        ("TRANSACTIONS SHEET", "", LIGHT_BLUE),
        ("Column", "What to enter", BLUE),
        ("Account No", "Must match an Account No from Clients sheet", WHITE),
        ("Date", "Transaction date — any readable format", ACCENT),
        ("Description", "Transaction narration / description", WHITE),
        ("Reference", "UTR / NEFT / UPI reference (optional)", ACCENT),
        ("Debit", "Amount going out (leave blank if credit)", WHITE),
        ("Credit", "Amount coming in (leave blank if debit)", ACCENT),
        ("Balance", "Running balance after this transaction", WHITE),
    ]
    for r, (col_b, col_c, bg) in enumerate(rows, 1):
        ws_i.row_dimensions[r].height = 18
        for c in range(1, 8):
            ws_i.cell(r, c).fill = PatternFill("solid", fgColor=bg)
        cell_b = ws_i.cell(r, 2, value=col_b)
        cell_c = ws_i.cell(r, 3, value=col_c)
        bold = bg in (BLUE, LIGHT_BLUE)
        fg = WHITE if bold else "000000"
        cell_b.font = Font(name="Arial", bold=bold, color=fg, size=10)
        cell_c.font = Font(name="Arial", color=fg, size=9)
        cell_b.border = cell_c.border = BORDER
        cell_b.alignment = Alignment(vertical="center")
        cell_c.alignment = Alignment(vertical="center")

    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    path = make_sample_excel(Path("sample_client_data.xlsx"))
    print(f"Sample Excel created: {path}")
